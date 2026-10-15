# בונה קובץ ZIP עם הפרומפטים הקבועים בקבצי Word (חלקים משתנים מסומנים בצבע) ובטקסט נקי.
# הרצה (מתיקיית sim-prompt-builder):
#   node tools/export-prompts.js > tools/out/prompts.json
#   python tools/make_zip.py
import io
import json
import os
import re
import shutil
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

ROOT = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(ROOT, 'out')
SITE_URL = 'https://orlevs.github.io/sim-prompt-builder/'
ZIP_NAME = 'sim-prompts.zip'

S0, S1 = '', ''
FONT = 'Arial'

COLORS = {
    'fill': 'FFE58F',    # צהוב: ממלאים לכל תרגיל או תלמיד/ה
    'adjust': 'BFDDF7',  # כחול: משתנה לפי ההגדרות
    'group': 'C9EBC4',   # ירוק: מוסיפים רק כשיש יותר מתלמיד/ה אחד/ת
}

# ------------------------------------------------------------- RTL helpers
_HEB = re.compile(r'[֐-׿יִ-ﭏ]')


def _strong(ch):
    if _HEB.match(ch):
        return True
    if ch.isascii() and ch.isalnum():
        return False
    return None


def _split_by_script(text):
    default_rtl = next((s for s in (_strong(c) for c in text) if s is not None), True)
    segs, buf, buf_rtl = [], '', None
    for ch in text:
        s = _strong(ch)
        kind = s if s is not None else (buf_rtl if buf_rtl is not None else default_rtl)
        if buf_rtl is None or kind == buf_rtl:
            buf, buf_rtl = buf + ch, kind
        else:
            segs.append((buf, buf_rtl))
            buf, buf_rtl = ch, kind
    if buf:
        segs.append((buf, buf_rtl))
    out = [[s, r] for s, r in segs]
    for i in range(len(out) - 1):
        if out[i][1] is False and out[i + 1][1] is True and out[i][0].endswith(' '):
            st = out[i][0].rstrip(' ')
            out[i + 1][0] = out[i][0][len(st):] + out[i + 1][0]
            out[i][0] = st
    return [(s, r) for s, r in out if s]


_LTR_ISLAND = re.compile(r'[A-Za-z0-9](?:[^֐-׿\n]*[A-Za-z0-9])?')


def _ltr_islands(text):
    """מחלק לריצות: רצפים של אנגלית/מספרים (כולל מה שביניהם, בלי עברית) הם LTR; כל השאר RTL."""
    out, pos = [], 0
    for m in _LTR_ISLAND.finditer(text):
        if m.start() > pos:
            out.append((text[pos:m.start()], True))
        out.append((m.group(0), False))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], True))
    return out


def _para_is_rtl(text):
    if _HEB.search(text):
        return True
    if any(ch.isascii() and ch.isalnum() for ch in text):
        return False
    return True


def parse_marked(line, default_kind=None):
    """שורה עם סימוני S0…S1 לרשימת (טקסט, סוג)."""
    parts, kind, buf = [], default_kind, ''
    for ch in line:
        if ch == S0:
            if buf:
                parts.append((buf, kind))
            buf, kind = '', 'var'
        elif ch == S1:
            if buf:
                parts.append((buf, kind))
            buf, kind = '', default_kind
        else:
            buf += ch
    if buf:
        parts.append((buf, kind))
    res = []
    for t, k in parts:
        if k == 'var':
            k = 'fill' if t.strip().startswith('⟦') else 'adjust'
        res.append((t, k))
    return res


def _fill_runs(p, parts, size=11, bold=False, color=None):
    """parts: [(text, kind)] או [(text, kind, bold)]. kind: None/'fill'/'adjust'/'group'."""
    parts = [(x[0], x[1], x[2] if len(x) > 2 else False) for x in parts]
    full = ''.join(t for t, _, _ in parts)
    has_latin = any(ch.isascii() and ch.isalpha() for ch in full)
    for text, kind, pbold in parts:
        for seg, is_rtl in _ltr_islands(text):
            run = p.add_run(seg)
            rPr = run._r.get_or_add_rPr()
            rPr.append(rPr.makeelement(qn('w:rFonts'), {qn('w:ascii'): FONT, qn('w:hAnsi'): FONT, qn('w:cs'): FONT}))
            if bold or pbold:
                rPr.append(rPr.makeelement(qn('w:b'), {}))
                rPr.append(rPr.makeelement(qn('w:bCs'), {}))
            if color:
                rPr.append(rPr.makeelement(qn('w:color'), {qn('w:val'): color}))
            rPr.append(rPr.makeelement(qn('w:sz'), {qn('w:val'): str(int(size * 2))}))
            rPr.append(rPr.makeelement(qn('w:szCs'), {qn('w:val'): str(int(size * 2))}))
            if kind in COLORS:
                rPr.append(rPr.makeelement(qn('w:shd'), {qn('w:val'): 'clear', qn('w:color'): 'auto', qn('w:fill'): COLORS[kind]}))
            # עברית וסימני פיסוק: <w:rtl/>. רצף אנגלית/מספרים: ריצה נפרדת בלי הסימון, כדי ש-Word לא יהפוך אותו.
            if is_rtl:
                rPr.append(rPr.makeelement(qn('w:rtl'), {}))


def _rtl_para(p):
    # <w:bidi/> בלי יישור מפורש: פסקה מימין לשמאל מתחילה מימין.
    # (יישור "ימין" ב-Word על פסקה כזאת דווקא דוחף אותה שמאלה.)
    pPr = p._p.get_or_add_pPr()
    pPr.append(pPr.makeelement(qn('w:bidi'), {}))


def add_par(doc, parts, size=11, bold=False, space_after=4, color=None, space_before=0):
    p = doc.add_paragraph()
    _rtl_para(p)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.2
    _fill_runs(p, parts, size=size, bold=bold, color=color)
    return p


def T(doc, text, **kw):
    return add_par(doc, [(text, kw.pop('kind', None))], **kw)


def new_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(2)
    sec._sectPr.append(sec._sectPr.makeelement(qn('w:bidi'), {}))
    st = doc.styles['Normal']
    st.font.name = FONT
    st.font.size = Pt(11)
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = rpr.makeelement(qn('w:rFonts'), {})
        rpr.append(rfonts)
    rfonts.set(qn('w:cs'), FONT)
    return doc


def legend(doc):
    add_par(doc, [('מקרא: ', None), (' צהוב ', 'fill'), (' = ממלאים לכל תרגיל או תלמיד/ה.  ', None),
                  (' כחול ', 'adjust'), (' = משתנה לפי ההגדרות (לשון פנייה, תווית, שם).  ', None),
                  (' ירוק ', 'group'), (' = מוסיפים רק כשיש יותר מתלמיד/ה אחד/ת.', None)], size=10, space_after=6)


def md_parts(line, default_kind=None):
    """שורה עם סימוני משתנים ו-**מודגש** לרשימת (טקסט, סוג, מודגש)."""
    out = []
    bold = False
    for text, kind in parse_marked(line, default_kind):
        chunks = text.split('**')
        for i, c in enumerate(chunks):
            if i > 0:
                bold = not bold
            if c:
                out.append((c, kind, bold))
    return out


def add_table(doc, rows, default_kind=None):
    from docx.enum.table import WD_TABLE_DIRECTION, WD_TABLE_ALIGNMENT
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = 'Table Grid'
    t.table_direction = WD_TABLE_DIRECTION.RTL
    t.alignment = WD_TABLE_ALIGNMENT.RIGHT
    for i, r in enumerate(rows):
        for j in range(ncol):
            cell = t.rows[i].cells[j]
            p = cell.paragraphs[0]
            _rtl_para(p)
            p.paragraph_format.space_after = Pt(2)
            txt = r[j] if j < len(r) else ''
            _fill_runs(p, md_parts(txt, default_kind), size=10.5, bold=(i == 0))
            if i == 0:
                tcPr = cell._tc.get_or_add_tcPr()
                tcPr.append(tcPr.makeelement(qn('w:shd'), {qn('w:val'): 'clear', qn('w:color'): 'auto', qn('w:fill'): 'E8F2FB'}))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _cells(line):
    s = line.strip()
    if s.startswith('|'):
        s = s[1:]
    if s.endswith('|'):
        s = s[:-1]
    return [c.strip() for c in s.split('|')]


def write_prompt(doc, text, default_kind=None):
    """מציג את הפרומפט כמסמך קריא: כותרות, טבלאות ותבליטים במקום סימני Markdown."""
    lines = text.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i]
        plain = line.replace(S0, '').replace(S1, '').strip()
        if plain.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].replace(S0, '').replace(S1, '').strip().startswith('|'):
                raw = lines[i]
                if not re.fullmatch(r'[|\-\s:]+', raw.replace(S0, '').replace(S1, '').strip()):
                    rows.append(_cells(raw))
                i += 1
            add_table(doc, rows, default_kind)
            continue
        i += 1
        if not plain:
            continue
        m = re.match(r'^(#{1,3}) ', plain)
        if m:
            level = len(m.group(1))
            body = re.sub(r'^(?)#{1,3} ', r'\1', line.lstrip())
            size = {1: 16, 2: 13.5, 3: 12}[level]
            color = {1: '0B2545', 2: '1D5FA8', 3: '0B2545'}[level]
            add_par(doc, md_parts(body, default_kind), size=size, bold=True, color=color,
                    space_before=(10 if level < 3 else 6), space_after=3)
            continue
        stripped = line.lstrip()
        indent = len(line) - len(stripped)
        if re.match(r'^()?- ', stripped):
            body = re.sub(r'^(?)- ', r'\1', stripped)
            add_par(doc, [('    ' * (indent // 3) + '•  ', default_kind)] + md_parts(body, default_kind), space_after=2)
            continue
        add_par(doc, md_parts(line, default_kind), space_after=4)


def prompt_doc(path, title, where, notes, text):
    doc = new_doc()
    T(doc, title, size=18, bold=True, color='0B2545', space_after=4)
    T(doc, 'לאן מדביקים: ' + where, size=11, bold=True, color='1D5FA8', space_after=4)
    for n in notes:
        T(doc, '• ' + n, size=10.5, space_after=2)
    legend(doc)
    T(doc, '──────── הפרומפט מתחיל כאן ────────', size=10, color='667085', space_after=8)
    T(doc, 'בקובץ הזה הפרומפט מעוצב לקריאה. להעתקה למערכת: התיקייה "טקסט נקי להעתקה".', size=10, color='667085', space_after=8)
    write_prompt(doc, text)
    T(doc, '──────── סוף הפרומפט ────────', size=10, color='667085', space_after=0)
    doc.save(path)


def clean(text):
    return text.replace(S0, '').replace(S1, '')


# ------------------------------------------------------------- content
def build(d, stage):
    os.makedirs(stage, exist_ok=True)
    words = os.path.join(stage, 'Word - עם סימון צבעים')
    txts = os.path.join(stage, 'טקסט נקי להעתקה')
    pers_w = os.path.join(words, 'דמויות מהמאגר')
    pers_t = os.path.join(txts, 'דמויות מהמאגר')
    for p in (words, txts, pers_w, pers_t):
        os.makedirs(p, exist_ok=True)

    common_fill = 'כל מה שמסומן בצהוב כתוב כ־⟦שם השדה⟧: מוחקים את הסוגריים ומכניסים את התוכן.'
    files = [
        ('02 מנהל שיח - מעבר שלבים ידני', 'orchestrator_manual', 'פרומפט מנהל שיח',
         ['לבחור את הקובץ הזה כשהמורה מעביר/ה שלבים בעצמו/ה ([מעבר לשלב 2]). אחרת: קובץ 03.',
          common_fill,
          'בכחול: STUDENT_1 היא התווית שמנהל השיח מחזיר כשהתלמיד/ה מדבר/ת. לשנות לתווית שהמערכת מצפה לה (למשל ARBEL).',
          'יעדי השלבים: שורה לכל שלב בטבלה, בפורמט "| מספר. שם השלב | מה התלמיד/ה אומר/ת בעצמו/ה |".']),
        ('03 מנהל שיח - מעבר שלבים אוטומטי', 'orchestrator_auto', 'פרומפט מנהל שיח',
         ['לבחור את הקובץ הזה כשמנהל השיח מעביר שלבים בעצמו (NEXT_PART).',
          common_fill,
          'הרעיון המרכזי והפתרון לא חובה, אבל עוזרים למנהל השיח להחליט אם יעד הושג. אם לא ממלאים, מוחקים את כל המקטע "כדי לשפוט אם יעד הושג".']),
        ('04 מנחה - שומר על המסגרת', 'mentor', 'פרומפט מנחה',
         ['זהה בשני מצבי המעבר בין שלבים.',
          'המנחה לא פדגוגי: מדבר רק כשמנסים לשבור את הדמות, או כשהשפה לא מתאימה.',
          'הקובץ כתוב בלשון נקבה למורה. למורה בלשון זכר: להחליף את "בלשון נקבה" ל"בלשון זכר" (מסומן בכחול).',
          'בדוגמאות: לעדכן את שם התלמיד/ה ואת לשון הפנייה (כחול/צהוב).']),
        ('05 הנחיות משותפות לתלמידים - מעבר ידני', 'shared_manual', 'פרומפט על לכל תלמידי AI',
         ['מתאים לשיעור עם מעבר שלבים ידני. לאוטומטי: קובץ 06.',
          common_fill,
          'שורת המורה (כחול): לבחור "המורה מלמדת… לשון נקבה" או "המורה מלמד… לשון זכר".',
          'שלבי השיעור: שורה לכל שלב, "1. שם השלב". היעדים: טבלה כמו במנהל השיח.',
          'המקטע "עוד לא למדנו" לא חובה. אם אין, מוחקים אותו.']),
        ('06 הנחיות משותפות לתלמידים - מעבר אוטומטי', 'shared_auto', 'פרומפט על לכל תלמידי AI',
         ['מתאים לשיעור עם מעבר שלבים אוטומטי.', common_fill]),
        ('07 דמות - תבנית ריקה בלשון נקבה', 'character_f', 'פרומפט הדמות (אחד לכל תלמיד/ה)',
         ['לבניית דמות חדשה. לדמות מהמאגר: התיקייה "דמויות מהמאגר".',
          common_fill,
          'תגובות למורה: שורה לכל תגובה בפורמט "| מה המורה עושה | איך הדמות מגיבה |".',
          'משפט פתיחה: אם לא רוצים משפט קבוע, משאירים את הניסוח של "משפט אחד קצר…" כמו שהוא.']),
        ('08 דמות - תבנית ריקה בלשון זכר', 'character_m', 'פרומפט הדמות (אחד לכל תלמיד/ה)',
         ['כמו קובץ 07, בלשון זכר.', common_fill]),
        ('09 פרומפט עזר - הצעת טעות מ-AI', 'helper', 'לא נכנס לסימולציה. מדביקים בשיחה רגילה עם Claude',
         ['ממלאים את פרטי התרגיל והדמות, ומדביקים בשיחה חדשה עם Claude.',
          'מקבלים שלוש הצעות. את מה שמתחת ל"להדבקה" מעתיקים לסעיף "הטעות שלך" בפרומפט הדמות.']),
    ]
    for name, key, where, notes in files:
        title = name[3:]
        prompt_doc(os.path.join(words, name + '.docx'), title, where, notes, d[key])
        with io.open(os.path.join(txts, name + '.txt'), 'w', encoding='utf-8-sig', newline='\r\n') as f:
            f.write(clean(d[key]))

    for i, p in enumerate(d['personas'], 1):
        name = '%d %s' % (i, p['name'])
        gtxt = 'לשון נקבה' if p['g'] == 'f' else 'לשון זכר'
        prompt_doc(os.path.join(pers_w, name + '.docx'), 'דמות: ' + p['name'], 'פרומפט הדמות',
                   [p['tagline'], 'כתוב ב' + gtxt + '. לשון אחרת: לבנות באתר, שם זה מתחלף אוטומטית.',
                    'ממלאים בצהוב: הכיתה והטעות בתרגיל. השם והתווית (כחול) אפשר לשנות.'],
                   p['text'])
        with io.open(os.path.join(pers_t, name + '.txt'), 'w', encoding='utf-8-sig', newline='\r\n') as f:
            f.write(clean(p['text']))

    instructions(d, os.path.join(stage, '01 הוראות - מה מכניסים איפה.docx'))


def instructions(d, path):
    doc = new_doc()
    T(doc, 'הפרומפטים הקבועים לסימולציית הוראה', size=20, bold=True, color='0B2545', space_after=2)
    T(doc, 'מה יש בקובץ, ומה מכניסים איפה', size=13, color='1D5FA8', space_after=10)

    T(doc, 'הדרך הקלה', size=14, bold=True, space_after=2)
    T(doc, 'באתר ממלאים טופס ומקבלים את כל הפרומפטים מוכנים, בלי לערוך ידנית: ' + SITE_URL + ' (סיסמה: 2221).', space_after=8)
    T(doc, 'הקבצים כאן מיועדים למי שרוצה לראות את הבסיס, לערוך ידנית, או לעבוד בלי האתר.', space_after=10)

    T(doc, 'מקרא הצבעים', size=14, bold=True, space_after=2)
    add_par(doc, [(' צהוב ', 'fill'), (' — ממלאים לכל תרגיל או תלמיד/ה. כתוב כ־⟦שם השדה⟧: מוחקים את הסוגריים ומכניסים תוכן.', None)])
    add_par(doc, [(' כחול ', 'adjust'), (' — כבר כתוב, אבל משתנה לפי ההגדרות: לשון פנייה, שם, תווית.', None)])
    add_par(doc, [(' ירוק ', 'group'), (' — מוסיפים רק כשיש יותר מתלמיד/ה אחד/ת (בסוף המסמך הזה).', None)], space_after=4)
    T(doc, 'כל השאר קבוע, ולא משנים אותו.', space_after=10)

    T(doc, 'מה מדביקים בכל תיבה במערכת', size=14, bold=True, space_after=2)
    rows = [
        ('פרומפט מנהל שיח', 'קובץ 02 (מעבר שלבים ידני) או 03 (אוטומטי)'),
        ('פרומפט מנחה', 'קובץ 04'),
        ('פרומפט על לכל תלמידי AI', 'קובץ 05 (ידני) או 06 (אוטומטי)'),
        ('פרומפט דמות (אחד לכל תלמיד/ה)', 'דמות מהמאגר, או תבנית 07 או 08'),
    ]
    for a, b in rows:
        add_par(doc, [('• ' + a + ': ', None), (b, None)])
    T(doc, 'קובץ 09 לא נכנס למערכת. הוא לשיחה רגילה עם Claude, כדי לקבל הצעות לטעות של הדמות.', space_after=10)

    T(doc, 'סדר עבודה', size=14, bold=True, space_after=2)
    steps = [
        'מחליטים: מעבר שלבים ידני או אוטומטי. זה קובע איזה קובץ לוקחים למנהל השיח ולהנחיות המשותפות.',
        'ממלאים את פרטי התרגיל בהנחיות המשותפות: שם, כיתה, מה כבר למדו, המשימה, שרטוט, ומה עוד לא למדו.',
        'כותבים את שלבי השיעור, ולכל שלב יעד: מה התלמיד/ה אומר/ת בעצמו/ה בסוף השלב. אותם יעדים נכנסים גם למנהל השיח.',
        'בוחרים דמות מהמאגר או ממלאים תבנית. כותבים את הטעות שלה בתרגיל (או מבקשים הצעה עם קובץ 09).',
        'במנהל השיח ובמנחה: מעדכנים שם, תווית ולשון פנייה של כל תלמיד/ה.',
        'מעתיקים מתיקיית "טקסט נקי להעתקה" כדי שלא ייכנסו צבעים או עיצוב למערכת.',
    ]
    for i, s in enumerate(steps, 1):
        T(doc, '%d. %s' % (i, s), space_after=3)
    T(doc, '', space_after=6)

    T(doc, 'השדות הצהובים', size=14, bold=True, space_after=2)
    fields = [
        ('⟦שם התרגיל⟧', 'למשל: הפקעה ופיצוי 1, סעיף 3'),
        ('⟦כיתה ורמה⟧', 'למשל: כיתה ט׳, הקבצה א׳'),
        ('⟦מה כבר למדו⟧', 'שורה לכל דבר, וכל שורה מתחילה במקף'),
        ('⟦טקסט המשימה⟧', 'המשימה כפי שהתלמידים רואים אותה, כולל השאלה'),
        ('⟦שלבי השיעור⟧', 'שורה לכל שלב: מספר, נקודה ושם השלב'),
        ('⟦שלבים ויעדים⟧', 'טבלה עם שתי עמודות: שם השלב, ומה התלמיד/ה אומר/ת בעצמו/ה'),
        ('⟦שם⟧ / ⟦שם התלמיד/ה⟧', 'שם הדמות'),
        ('⟦הטעות של הדמות⟧', 'בגוף שלישי: מה הדמות עשתה ומה היא חושבת. בלי מה יתקן את זה'),
        ('⟦הרעיון המתמטי המרכזי⟧, ⟦הפתרון⟧', 'רק במעבר אוטומטי ובפרומפט העזר'),
    ]
    for a, b in fields:
        add_par(doc, [(a, 'fill'), (' — ' + b, None)], space_after=3)
    T(doc, '', space_after=6)

    T(doc, 'לשון פנייה', size=14, bold=True, space_after=2)
    T(doc, 'ההנחיות המשותפות כתובות עם לוכסנים (את/ה), כי הן לכל התלמידים. פרומפט הדמות כתוב בלשון אחת: לקחת את התבנית המתאימה (07 נקבה, 08 זכר). שורת המורה, וההערות של המנחה, מסומנות בכחול.', space_after=10)

    T(doc, 'כשיש יותר מתלמיד/ה אחד/ת', size=14, bold=True, space_after=2)
    T(doc, '1. בהנחיות המשותפות, בשורה הראשונה של "השיעור": במקום "שיעור פרטני…" כותבים:', space_after=2)
    add_par(doc, [('- בשיעור משתתפים: ', 'group'), ('⟦שמות התלמידים⟧', 'fill'), ('. אין תלמידים אחרים מלבדם, ולא מזכירים או מניחים שיש.', 'group')], space_after=6)
    T(doc, '2. בהנחיות המשותפות, לפני "איך כותבים", מוסיפים:', space_after=2)
    write_prompt(doc, d['group_shared'], default_kind='group')
    T(doc, '', space_after=4)
    T(doc, '3. במנהל השיח: בטבלת המשתתפים מוסיפים שורה לכל תלמיד/ה, וב"מה מחזירים" מוסיפים את כל התוויות. לפני "מתי MENTOR" מוסיפים:', space_after=2)
    write_prompt(doc, d['group_orch'], default_kind='group')
    T(doc, '', space_after=4)
    T(doc, '4. במנחה: בשורה "התלמידים בשיעור" מוסיפים את כל התלמידים עם לשון הפנייה של כל אחד/ת.', space_after=2)
    T(doc, '5. פרומפט דמות נפרד לכל תלמיד/ה.', space_after=10)

    T(doc, 'מה יש בתיקיות', size=14, bold=True, space_after=2)
    T(doc, '• Word - עם סימון צבעים: לקריאה ולעריכה.', space_after=2)
    T(doc, '• טקסט נקי להעתקה: אותו תוכן בלי צבעים, להעתקה למערכת.', space_after=2)
    T(doc, '• דמויות מהמאגר: שמונה פרומפטי דמות מוכנים. נשאר למלא כיתה וטעות.', space_after=2)
    doc.save(path)


def main():
    with io.open(os.path.join(OUT, 'prompts.json'), encoding='utf-8') as f:
        d = json.load(f)
    import tempfile
    work = tempfile.mkdtemp(prefix='simzip-')
    stage = os.path.join(work, 'פרומפטים קבועים - סימולציית הוראה')
    if os.path.isdir(stage):
        shutil.rmtree(stage)
    build(d, stage)
    site_dl = os.path.join(ROOT, '..', 'downloads')
    os.makedirs(site_dl, exist_ok=True)
    zpath = os.path.join(site_dl, ZIP_NAME)
    with zipfile.ZipFile(zpath, 'w', zipfile.ZIP_DEFLATED) as z:
        for base, _, names in os.walk(stage):
            for n in sorted(names):
                full = os.path.join(base, n)
                z.write(full, os.path.relpath(full, work))
    print(zpath, os.path.getsize(zpath))


if __name__ == '__main__':
    main()
